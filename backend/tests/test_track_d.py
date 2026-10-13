from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException
from pptx import Presentation
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.evaluation.evaluator import reference_token_f1, run_model_evaluation
from app.ingestion.ppt_processor import extract_ppt_chunks
from app.ingestion.processor import (
    assign_explicit_prerequisites,
    classify_topic_and_concept,
    detect_file_type,
)
from app.ingestion.video_processor import _format_timestamp, extract_video_chunks, whisper
from app.learner.mastery import get_mastery, get_weak_topics
from app.learner.misconceptions import detect_misconception
from app.learner.recommendations import generate_recommendations
from app import main
from app.llm import LLMError
from app.main import _citations_match_retrieved_chunks, submit_assessment
from app.models import (
    AssessmentLog,
    AssessmentQuestion,
    Concept,
    ContentChunk,
    Material,
    Topic,
)
from app.rag.vector_store import _simple_embedding
from app.schemas import AskResponse, AssessmentSubmission, Citation


class TrackDTests(unittest.TestCase):
    def setUp(self) -> None:
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.session = sessionmaker(bind=self.engine)()

    def tearDown(self) -> None:
        self.session.close()
        self.engine.dispose()

    def test_citations_must_match_retrieved_source_and_location(self) -> None:
        chunks = [{"source_name": "DBMS course.pdf", "location_ref": "Page 4"}]
        self.assertTrue(
            _citations_match_retrieved_chunks(
                "Durability is defined here. [DBMS course.pdf | Page 4]", chunks
            )
        )
        self.assertFalse(
            _citations_match_retrieved_chunks(
                "See OtherBook.pdf, Page 99.", chunks
            )
        )
        self.assertFalse(_citations_match_retrieved_chunks("No citation included.", chunks))

    def test_fallback_embedding_is_stable_across_processes(self) -> None:
        backend_dir = Path(__file__).resolve().parents[1]
        code = (
            "import json; from app.rag.vector_store import _simple_embedding; "
            "print(json.dumps(_simple_embedding('database transaction durability').tolist()))"
        )
        output = subprocess.check_output(
            [sys.executable, "-c", code], cwd=backend_dir, text=True
        )
        self.assertEqual(
            json.loads(output), _simple_embedding("database transaction durability").tolist()
        )

    def test_pptx_extraction_and_topic_classification(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "dbms.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Transactions"
            slide.placeholders[1].text = "ACID durability protects committed changes."
            presentation.save(path)

            chunks = extract_ppt_chunks(path, path.name)

        self.assertEqual(chunks[0]["location_ref"], "Slide 1")
        self.assertEqual(chunks[0]["source_name"], "dbms.pptx")
        self.assertEqual(
            classify_topic_and_concept(chunks[0]["content"])[0],
            "Transactions and Concurrency",
        )
        self.assertEqual(detect_file_type("course.pptx"), ".pptx")
        with self.assertRaises(ValueError):
            detect_file_type("legacy.ppt")

    def test_prerequisite_requires_explicit_course_statement(self) -> None:
        material = Material(filename="prerequisite-source.pdf", file_type="pdf")
        self.session.add(material)
        self.session.flush()
        source_text = "ACID is a prerequisite for serializability."
        self.session.add(
            ContentChunk(
                material_id=material.id,
                content=source_text,
                source_name=material.filename,
                source_type="pdf",
                location_ref="Page 7",
                chunk_hash="explicit-prerequisite-test",
            )
        )
        self.session.flush()

        assign_explicit_prerequisites(self.session, source_text)
        self.session.flush()
        acid = self.session.query(Concept).filter_by(name="Acid").one()
        serializability = self.session.query(Concept).filter_by(name="Serializability").one()
        self.assertEqual(serializability.prerequisite_id, acid.id)

        assign_explicit_prerequisites(self.session, "ACID helps with serializability.")
        assign_explicit_prerequisites(
            self.session, "Serializability is a prerequisite for ACID."
        )
        self.session.flush()
        self.assertEqual(serializability.prerequisite_id, acid.id)
        self.assertIsNone(acid.prerequisite_id)

    def test_submission_updates_feedback_mastery_weak_topics_and_recommendation(self) -> None:
        topic = Topic(name="Transactions and Concurrency")
        material = Material(filename="course.pdf", file_type="pdf")
        self.session.add_all([topic, material])
        self.session.flush()
        concept = Concept(topic_id=topic.id, name="Durability")
        self.session.add(concept)
        self.session.flush()
        self.session.add(
            ContentChunk(
                material_id=material.id,
                content="Durability preserves committed changes after failure.",
                source_name=material.filename,
                source_type="pdf",
                location_ref="Page 4",
                topic_id=topic.id,
                concept_id=concept.id,
                chunk_hash="submission-flow-test",
            )
        )
        question = AssessmentQuestion(
            topic=topic.name,
            concept=concept.name,
            difficulty="medium",
            question_type="mcq",
            question_text="Which ACID property protects committed changes?",
            options=["Atomicity", "Durability"],
            correct_answer="Durability",
            source_reference="course.pdf | Page 4",
        )
        self.session.add(question)
        self.session.commit()

        correct = submit_assessment(
            AssessmentSubmission(question_id=question.id, student_answer="Durability"),
            self.session,
        )
        mastery_after_correct = get_mastery(self.session)[0]["mastery_score"]
        incorrect = submit_assessment(
            AssessmentSubmission(question_id=question.id, student_answer="Atomicity"),
            self.session,
        )
        mastery_after_incorrect = get_mastery(self.session)[0]["mastery_score"]

        self.assertTrue(correct.correctness)
        self.assertEqual(correct.citations[0], Citation(source_name="course.pdf", location="Page 4"))
        self.assertFalse(incorrect.correctness)
        self.assertEqual(incorrect.score, 0.0)
        self.assertIn("not supported", incorrect.source_grounded_feedback)
        self.assertIn("durability with atomicity", incorrect.misconception_detected)
        self.assertGreater(mastery_after_correct, mastery_after_incorrect)
        self.assertEqual(get_weak_topics(self.session)[0]["topic"], topic.name)
        recommendation = generate_recommendations(self.session)[0]
        self.assertEqual(recommendation["topic"], topic.name)
        self.assertEqual(recommendation["source_name"], material.filename)
        self.assertEqual(recommendation["location"], "Page 4")
        logs = self.session.query(AssessmentLog).order_by(AssessmentLog.id).all()
        self.assertEqual(len(logs), 2)
        self.assertIsNotNone(logs[1].misconception_detected)

    def test_live_evaluation_uses_supplied_model_output_without_mastery_mutation(self) -> None:
        question = AssessmentQuestion(
            topic="Transactions and Concurrency",
            concept="Durability",
            difficulty="medium",
            question_type="mcq",
            question_text="What does durability guarantee?",
            options=["Committed changes persist", "Transactions are isolated"],
            correct_answer="Committed changes persist",
            source_reference="course.pdf | Page 4",
        )
        self.session.add(question)
        self.session.commit()
        chunks = [
            {
                "content": "Committed changes persist after system failure.",
                "source_name": "course.pdf",
                "location_ref": "Page 4",
            }
        ]

        def answer_with_sources(question_text: str, topic: str, concept: str):
            answer = "Committed changes persist. [course.pdf | Page 4]"
            return AskResponse(
                answer=answer,
                supported=True,
                citations=[Citation(source_name="course.pdf", location="Page 4")],
            ), chunks

        result = run_model_evaluation(
            self.session, answer_with_sources, _citations_match_retrieved_chunks
        )
        self.assertEqual(result["cases_evaluated"], 1)
        self.assertEqual(result["supported_response_rate"], 1.0)
        self.assertEqual(result["citation_validity"], 1.0)
        self.assertGreater(result["mean_reference_token_f1"], 0.0)
        self.assertEqual(get_mastery(self.session), [])

    def test_timestamp_format_and_missing_whisper_are_explicit(self) -> None:
        self.assertEqual(
            [_format_timestamp(value) for value in (0, 59.9, 60, 3661.8)],
            ["00:00", "00:59", "01:00", "01:01:01"],
        )
        if whisper is None:
            with tempfile.TemporaryDirectory() as directory:
                video = Path(directory) / "untranscribed.mp4"
                video.write_bytes(b"")
                with self.assertRaisesRegex(RuntimeError, "Whisper is not installed"):
                    extract_video_chunks(video, video.name)

    def test_live_evaluation_provider_failure_does_not_store_metrics(self) -> None:
        self.session.add(
            AssessmentQuestion(
                topic="Transactions and Concurrency",
                concept="Durability",
                difficulty="medium",
                question_type="mcq",
                question_text="What does durability guarantee?",
                options=["Committed changes persist", "Transactions are isolated"],
                correct_answer="Committed changes persist",
                source_reference="course.pdf | Page 4",
            )
        )
        self.session.commit()
        main.app.state.evaluation_results = None

        with patch.object(main, "_answer_with_sources", side_effect=LLMError("provider unavailable")):
            with self.assertRaises(HTTPException) as raised:
                main.run_evaluation(self.session)

        self.assertEqual(raised.exception.status_code, 503)
        self.assertIsNone(main.app.state.evaluation_results)

    def test_reference_token_f1_is_transparent(self) -> None:
        self.assertEqual(
            reference_token_f1(
                "Durability preserves committed changes",
                "Durability preserves changes",
            ),
            0.8571,
        )


if __name__ == "__main__":
    unittest.main()